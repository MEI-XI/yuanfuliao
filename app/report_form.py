"""Fill the supplied plain inspection form without rebuilding its table or signatures."""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


HEADERS = ('检验项目', '标准规定', '检验结果', '判定结论')


def _align_grid(table):
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for no_wrap in cell._tc.xpath('./w:tcPr/w:noWrap'):
                no_wrap.getparent().remove(no_wrap)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _align_report_header(table):
    for row in table.rows:
        if row.cells[0].text.strip().startswith('检验项目'):
            break
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def align_inspection_tables(source):
    document = Document(BytesIO(source))
    def visit(tables):
        for table in tables:
            if table.rows and any('检品名称' in cell.text for cell in table.rows[0].cells):
                _align_report_header(table)
            if table.rows and tuple(cell.text.strip() for cell in table.rows[0].cells) == HEADERS:
                _align_grid(table)
            for row in table.rows:
                for cell in row.cells:
                    visit(cell.tables)
    visit(document.tables)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def fill_plain_report(source, context):
    document = Document(BytesIO(source))
    xml = document.element.xml
    if '{{' in xml or '{%' in xml or not document.tables:
        return None
    table = document.tables[0]
    if (len(table.rows) != 7 or len(table.columns) != 4
            or '检品名称' not in table.cell(0, 0).text
            or '检验项目' not in table.cell(6, 0).text):
        return None
    fields = {(0, 1): 'material_name', (0, 3): 'company_name', (1, 1): 'batch_blank',
              (1, 3): 'specification', (2, 1): 'packaging', (2, 3): 'received_quantity',
              (3, 1): 'basis', (3, 3): 'production_dates', (4, 1): 'shelf_life', (4, 3): 'inspection_date'}
    for (row, col), key in fields.items():
        paragraph = table.cell(row, col).paragraphs[0]
        paragraph.clear().add_run(context.get(key) or '')
    table.cell(5, 0).paragraphs[0].clear().add_run('该公司是否提供内检报告：' + (context.get('internal_report') or ''))
    table.cell(5, 2).paragraphs[0].clear().add_run(context.get('report_reference') or '')
    _align_report_header(table)
    cell = table.cell(6, 0)
    paragraphs = cell.paragraphs
    if len(paragraphs) < 7:
        return None
    if table._tbl.tblPr.find(qn('w:tblpPr')) is not None:
        # A title following a floating table flows to its last page when the table grows.
        for sibling in table._tbl.itersiblings():
            if sibling.tag == qn('w:p'):
                text = ''.join(node.text or '' for node in sibling.xpath('.//w:t'))
                if ''.join(text.split()) == '原辅料检验报告单':
                    table._tbl.addprevious(sibling)
                    break
    conclusion = paragraphs[5]
    # A real four-column grid keeps wrapped values under their own headers.
    # Only the five old text paragraphs are replaced; anchored signatures below stay intact.
    grid = cell.add_table(rows=max(4, len(context['items'])) + 1, cols=4)
    grid.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    grid._tbl.tblPr.append(borders)
    width = cell.width or Pt(500)
    for col, proportion in zip(grid.columns, (.23, .30, .27, .20)):
        col.width = int(width * proportion)
        for item_cell in col.cells:
            item_cell.width = col.width
    for index, row in enumerate(grid.rows):
        row.height = Pt(22)
        item = context['items'][index - 1] if 0 < index <= len(context['items']) else {}
        values = HEADERS if index == 0 else tuple(item.get(key, '') for key in ('name', 'standard', 'result', 'conclusion'))
        for item_cell, value in zip(row.cells, values):
            paragraph = item_cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(value or '')
            run.font.size = Pt(12)
    _align_grid(grid)
    conclusion._p.addprevious(grid._tbl)
    for paragraph in paragraphs[:5]:
        paragraph._p.getparent().remove(paragraph._p)
    # Clearing just text preserves any anchored objects in the conclusion too.
    for node in conclusion._p.xpath('.//w:t'):
        node.text = ''
    conclusion.add_run('结  论：' + (context.get('conclusion') or ''))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
