"""Normalize a single uploaded template without retaining it on the server."""
import base64
import os
import shutil
import subprocess
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from docx import Document
from fastapi import HTTPException
from lxml import etree
from openpyxl import load_workbook

from app.personal_templates import decode_template

EXTENSIONS = {
    'docx': {'doc', 'docx', 'docm', 'dot', 'dotx', 'dotm', 'rtf', 'odt', 'xml'},
    'xlsx': {'xls', 'xlsx', 'xlsm', 'xlsb', 'xlt', 'xltx', 'xltm', 'ods', 'csv', 'xml'},
}
OOXML = {'docx', 'docm', 'dotx', 'dotm', 'xlsx', 'xlsm', 'xltx', 'xltm'}
CONTENT_TYPES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
}


def _package(data, kind):
    """Check expansion limits before an Office reader expands the package."""
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 3000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
            raise ValueError('模板解压后过大')
        if ('word/document.xml' if kind == 'docx' else 'xl/workbook.xml') not in archive.namelist():
            raise ValueError('文件内容与所选模板类型不符')


def _normalize_package(data, kind):
    _package(data, kind)
    output = BytesIO()
    if kind == 'xlsx':
        workbook = load_workbook(BytesIO(data))
        workbook.template = False
        workbook.save(output)
        workbook.close()
    else:
        # python-docx expects the document main content type. Remove VBA parts
        # and their relationships before opening macro/template variants.
        package = BytesIO()
        with ZipFile(BytesIO(data)) as source, ZipFile(package, 'w', ZIP_DEFLATED) as target:
            removed = {name for name in source.namelist() if 'vba' in name.lower()}
            for name in source.namelist():
                if name in removed:
                    continue
                content = source.read(name)
                if name == '[Content_Types].xml' or name.endswith('.rels'):
                    root = etree.fromstring(content, etree.XMLParser(resolve_entities=False, no_network=True))
                    for element in list(root):
                        if 'vba' in (element.get('PartName', '') + element.get('Type', '') + element.get('Target', '')).lower():
                            root.remove(element)
                        elif element.get('PartName') == '/word/document.xml':
                            element.set('ContentType', CONTENT_TYPES[kind])
                    content = etree.tostring(root)
                target.writestr(name, content)
        Document(package).save(output)
    return output.getvalue()


def _validate_legacy(data, extension, kind):
    ole = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
    if extension in {'doc', 'dot', 'xls', 'xlt'}:
        valid = data.startswith(ole)
    elif extension == 'rtf':
        valid = data.lstrip().startswith(b'{\\rtf')
    elif extension in {'odt', 'ods', 'xlsb'}:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                raise ValueError('expanded size')
            if extension == 'xlsb':
                valid = 'xl/workbook.bin' in archive.namelist()
            else:
                valid = archive.read('mimetype').strip() == ('application/vnd.oasis.opendocument.' + ('text' if extension == 'odt' else 'spreadsheet')).encode()
    elif extension == 'xml':
        root = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
        namespace = etree.QName(root).namespace or ''
        valid = namespace in ({'urn:schemas-microsoft-com:office:word', 'http://schemas.microsoft.com/office/2006/xmlPackage'}
                              if kind == 'docx' else {'urn:schemas-microsoft-com:office:spreadsheet'})
    else:  # CSV is a text spreadsheet; Office handles its delimiter/encoding.
        valid = not data.startswith((b'PK', ole)) and b'\x00' not in data
        try:
            data.decode('utf-8-sig')
        except UnicodeDecodeError:
            data.decode('gb18030')
    if not valid:
        raise ValueError('type')


def _libreoffice():
    configured = os.getenv('LIBREOFFICE_PATH')
    if configured:
        return configured if Path(configured).is_file() else None
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if executable:
        return executable
    for root in (os.getenv('ProgramFiles'), os.getenv('ProgramFiles(x86)')):
        if root:
            candidate = Path(root) / 'LibreOffice/program/soffice.exe'
            if candidate.is_file():
                return str(candidate)
    return None


def _office_available(kind):
    if os.name != 'nt':
        return False
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, ('Word' if kind == 'docx' else 'Excel') + '.Application\\CLSID'):
            return True
    except OSError:
        return False


def _run_converter(command, directory, timeout=45, environment=None):
    try:
        return subprocess.run(command, capture_output=True, timeout=timeout, check=False, env=environment,
                              creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    except subprocess.TimeoutExpired:
        # Only terminate the Office instance recorded by our own COM process.
        # Pre-existing Office process IDs are excluded by the script.
        pid_file = directory / 'office.pid'
        if os.name == 'nt' and pid_file.exists():
            own_pid = pid_file.read_text().strip()
            if own_pid.isdigit():
                subprocess.run(['taskkill', '/PID', own_pid, '/F'], capture_output=True,
                               timeout=5, creationflags=subprocess.CREATE_NO_WINDOW)
        raise HTTPException(504, '模板转换超时，请关闭该文件后重试，或另存为docx/xlsx再上传') from None


def _office_command(source, destination, kind, directory):
    # Paths are generated internally, passed as environment variables, never
    # interpolated into PowerShell source. New COM instances open only our copy.
    script = r'''
$ErrorActionPreference = 'Stop'
$app = $null
$document = $null
$existing = @(Get-Process -Name WINWORD,EXCEL -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
Add-Type -TypeDefinition 'using System; using System.Runtime.InteropServices; public static class WindowOwner { [DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr window, out uint process); }'
try {
    $app = New-Object -ComObject $env:TEMPLATE_OFFICE_APP
    $app.Visible = $false
    $app.AutomationSecurity = 3
    [uint32]$officeProcess = 0
    if ($env:TEMPLATE_KIND -eq 'xlsx') {
        [void][WindowOwner]::GetWindowThreadProcessId([IntPtr]$app.Hwnd, [ref]$officeProcess)
    }
    if ($officeProcess -and $existing -notcontains $officeProcess) {
        [IO.File]::WriteAllText($env:TEMPLATE_PID, [string]$officeProcess)
    }
    if ($env:TEMPLATE_KIND -eq 'docx') {
        $app.DisplayAlerts = 0
        $app.Options.UpdateLinksAtOpen = $false
        $document = $app.Documents.Open($env:TEMPLATE_SOURCE, $false, $true, $false, 'template-no-password')
        [void][WindowOwner]::GetWindowThreadProcessId([IntPtr]$document.ActiveWindow.Hwnd, [ref]$officeProcess)
        if ($officeProcess -and $existing -notcontains $officeProcess) {
            [IO.File]::WriteAllText($env:TEMPLATE_PID, [string]$officeProcess)
        }
        $document.SaveAs2($env:TEMPLATE_DESTINATION, 16)
    } else {
        $app.DisplayAlerts = $false
        $app.EnableEvents = $false
        $app.AskToUpdateLinks = $false
        $document = $app.Workbooks.Open($env:TEMPLATE_SOURCE, 0, $true, [Type]::Missing, 'template-no-password', [Type]::Missing, $true)
        $document.SaveAs($env:TEMPLATE_DESTINATION, 51)
    }
} finally {
    try {
        if ($null -ne $document) { $document.Close($false); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) }
    } finally {
        if ($null -ne $app) { $app.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($app) }
    }
}
'''
    # Use a child-specific environment, so simultaneous conversions cannot share
    # input/output paths or mutate the web process environment.
    environment = dict(os.environ, TEMPLATE_OFFICE_APP='Word.Application' if kind == 'docx' else 'Excel.Application',
                       TEMPLATE_KIND=kind, TEMPLATE_SOURCE=str(source), TEMPLATE_DESTINATION=str(destination),
                       TEMPLATE_PID=str(directory / 'office.pid'))
    encoded = base64.b64encode(script.encode('utf-16le')).decode('ascii')
    command = ['powershell.exe', '-NoProfile', '-NonInteractive', '-WindowStyle', 'Hidden', '-EncodedCommand', encoded]
    return command, environment


def _convert_legacy(data, extension, kind):
    with TemporaryDirectory(prefix='template-convert-') as temporary:
        directory = Path(temporary)
        source = directory / ('source.' + extension)
        destination = directory / ('source.' + kind)
        source.write_bytes(data)
        executable = _libreoffice()
        if executable:
            filter_name = 'docx:Office Open XML Text' if kind == 'docx' else 'xlsx:Calc MS Excel 2007 XML'
            _run_converter([executable, '-env:UserInstallation=' + (directory / 'profile').as_uri(), '--headless',
                            '--convert-to', filter_name, '--outdir', str(directory), str(source)], directory)
        office_available = _office_available(kind)
        if not destination.exists() and office_available:
            command, environment = _office_command(source, destination, kind, directory)
            # Pass paths only to this subprocess; restore nothing in global state.
            _run_converter(command, directory, environment=environment)
        if not destination.exists():
            if not executable and not office_available:
                raise HTTPException(503, '服务器未安装格式转换工具，请配置LibreOffice或另存为docx/xlsx后上传')
            raise HTTPException(422, '无法转换该模板，请确认文档未加密且能在Office中正常打开，或另存为docx/xlsx后上传')
        return destination.read_bytes()


def convert_template(data: bytes, filename: str, kind: str) -> bytes:
    extension = Path(filename).suffix.lower().lstrip('.')
    if extension not in EXTENSIONS[kind]:
        raise HTTPException(422, '请选择支持的Word或Excel文档，文件扩展名须与模板类型一致')
    if not data:
        raise HTTPException(422, '模板文件为空')
    try:
        if extension in OOXML:
            _package(data, kind)
            if extension != kind:
                data = _normalize_package(data, kind)
        else:
            _validate_legacy(data, extension, kind)
            data = _convert_legacy(data, extension, kind)
        return decode_template(base64.b64encode(data).decode('ascii'), kind)
    except (BadZipFile, ValueError, KeyError, OSError, etree.XMLSyntaxError) as exc:
        raise HTTPException(422, '模板文件损坏或内容与扩展名不符，请在Office中打开并另存后重试') from exc
