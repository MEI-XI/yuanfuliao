"""Read a personal template for one export; never write it to server templates."""
import base64
import binascii
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from lxml.etree import XMLSyntaxError

FILENAMES = {'xlsx': '查验台账.xlsx', 'docx': '检验报告.docx'}


def decode_template(encoded, kind):
    if encoded is None:
        return None
    try:
        data = base64.b64decode(encoded, validate=True)
        if not data or len(data) > 10 * 1024 * 1024:
            raise ValueError('size')
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                raise ValueError('expanded size')
            expected = 'word/document.xml' if kind == 'docx' else 'xl/workbook.xml'
            if expected not in archive.namelist():
                raise ValueError('type')
        if kind == 'xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data))
            sheet = workbook.active
            if not sheet or not isinstance(sheet['A1'].value, str) or '原辅料查验记录' not in sheet['A1'].value or sheet.max_column != 13:
                raise ValueError('layout')
            workbook.close()
        else:
            from docx import Document
            from docx.oxml.ns import qn
            document = Document(BytesIO(data))
            if document.element.tag != qn('w:document') or document.element.body is None:
                raise ValueError('document structure')
        return data
    except (binascii.Error, BadZipFile, ValueError, KeyError, OSError, XMLSyntaxError):
        raise HTTPException(422, '个人模板无法读取，请使用有效的检验报告.docx或保留标题及13列表头的查验台账.xlsx，文件不超过10MB') from None
