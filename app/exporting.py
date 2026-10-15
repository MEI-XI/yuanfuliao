import os
import re
import unicodedata
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from zoneinfo import ZoneInfo

from docxtpl import DocxTemplate
from jinja2.sandbox import SandboxedEnvironment

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def template_dir():
    return Path(os.getenv("TEMPLATE_DIR", "templates"))


def dates_text(values):
    return "\n".join(date.fromisoformat(value).strftime("%Y.%m.%d") for value in values)


def sourced(fields, key):
    value = fields.get(key, "")
    if value and fields.get(key + "_source") == "manufacturer" and not value.startswith("生产商"):
        return "生产商：" + value
    return value


def excel_bytes(records, ledger_names, template=None):
    workbook = load_workbook(BytesIO(template) if template is not None else template_dir() / "查验台账.xlsx")
    original = workbook.active
    sheets = []
    for ledger_id in dict.fromkeys(row.ledger_id for row in records):
        sheet = workbook.copy_worksheet(original)
        sheets.append((sheet, ledger_names[ledger_id]))
        company_records = [record for record in records if record.ledger_id == ledger_id]
        first_arrival = company_records[0].fields.get("arrival_date")
        year = f"{date.fromisoformat(first_arrival).year}年" if first_arrival else ""
        sheet['A1'] = re.sub(r"原辅料查验记录(?:\s*(?:X|\d{4})\s*年)?",
                             "原辅料查验记录" + year, sheet['A1'].value)
        # The supplied template has a title and header, no pre-filled business rows.
        if sheet.max_row > 2:
            sheet.delete_rows(3, sheet.max_row - 2)
        for index, record in enumerate(company_records, 3):
            f = record.fields
            values = [date.fromisoformat(f["arrival_date"]) if f.get("arrival_date") else None,
                      f.get("material_name"), dates_text(f.get("production_dates", [])).replace("\n", "、"), f.get("specification"),
                      f.get("company_name"), sourced(f, "address"), sourced(f, "contact"), f.get("shelf_life"),
                      f.get("received_quantity"), f.get("certificates"), f.get("sensory"), f.get("sampling_items"), f.get("notes")]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(index, column)
                cell._style = copy(sheet.cell(2, column)._style)
                cell.value = value or None
                if isinstance(value, str):
                    # Labels beginning '=' are text from a person/image, never formulas.
                    cell.data_type = "s"
            sheet.cell(index, 1).number_format = "yyyy.mm.dd"
        widths = {column: 10 for column in range(1, 14)}
        for row in sheet:
            row_height = 20
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                if cell.value is None:
                    continue
                text = cell.value.strftime('%Y.%m.%d') if isinstance(cell.value, date) else str(cell.value)
                lines = text.splitlines() or ['']
                row_height = max(row_height, len(lines) * max(16, (cell.font.sz or 11) * 1.3))
                # A merged title spans its columns and must not widen its first column.
                if any(cell.coordinate in merged and merged.max_col > merged.min_col for merged in sheet.merged_cells.ranges):
                    continue
                visible_width = max(sum(0 if unicodedata.combining(char) else 2 if unicodedata.east_asian_width(char) in 'WF' else 1
                                        for char in line) for line in lines)
                widths[cell.column] = max(widths.get(cell.column, 10), visible_width + 2)
            sheet.row_dimensions[row[0].row].height = min(409, row_height)
        for column, width in widths.items():
            sheet.column_dimensions[sheet.cell(2, column).column_letter].width = min(255, width)
        sheet.print_title_rows = "1:2"
        sheet.print_area = f"A1:M{sheet.max_row}"
    for template in list(workbook.worksheets):
        if all(template is not sheet for sheet, _ in sheets):
            workbook.remove(template)
    # Assign final names only after all template and temporary names are gone.
    import uuid
    for sheet, _ in sheets:
        sheet.title = uuid.uuid4().hex[:31]
    for sheet, name in sheets:
        sheet.title = name
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def word_bytes(record, template=None):
    source = template if template is not None else (template_dir() / "检验报告.docx").read_bytes()
    fields = record.fields
    report = record.report
    context = fields | report | {
        "production_dates": dates_text(fields.get("production_dates", [])),
        "inspection_date": dates_text([report["inspection_date"]]) if report.get("inspection_date") else "",
        "batch_blank": "", "items": [item | {'unit': ''} for item in record.items],
        "inspector": "", "approver": "",
        "report_reference": "其他检验项目依据该公司提供的内检报告为准" if report.get("internal_report") == "是" else "",
    }
    from app.report_form import fill_plain_report
    filled = fill_plain_report(source, context)
    if filled is not None:
        return filled
    template = DocxTemplate(BytesIO(source))
    template.render({key: "" if value is None else value for key, value in context.items()},
                    jinja_env=SandboxedEnvironment(autoescape=True), autoescape=True)
    output = BytesIO()
    template.save(output)
    return output.getvalue()


def report_filename(record):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', record.fields.get('material_name') or '').strip(' .') or '未命名'
    # Leave room for the date and duplicate suffix, even for surrogate-pair characters.
    name = name[:100]
    # Windows reserves device names even when they have an extension.
    if re.fullmatch(r'CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9]', name.split('.')[0], re.IGNORECASE):
        name = '_' + name
    inspection_date = record.report.get('inspection_date')
    report_date = date.fromisoformat(inspection_date) if inspection_date else record.created_at.astimezone(ZoneInfo('Asia/Shanghai')).date()
    return f'{name}_{report_date:%Y.%m.%d}.docx'


def word_archive(records, template=None):
    output = BytesIO()
    used_names = set()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for record in records:
            filename = report_filename(record)
            stem = filename[:-5]
            suffix = 0
            while filename.casefold() in used_names:
                suffix += 1
                filename = f'{stem}({suffix}).docx'
            used_names.add(filename.casefold())
            archive.writestr(filename, word_bytes(record, template))
    return output.getvalue()
