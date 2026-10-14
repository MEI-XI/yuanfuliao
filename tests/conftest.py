import os
import uuid

import psycopg
import pytest
from fastapi.testclient import TestClient
from pathlib import Path
from openpyxl import Workbook
from docx import Document

from app.main import create_app


@pytest.fixture
def client(tmp_path, monkeypatch):
    # Each test owns a schema. It never deletes or reads the user's business data.
    url = os.environ["TEST_DATABASE_URL"]
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    templates = tmp_path / "templates"
    templates.mkdir()
    monkeypatch.setenv("TEMPLATE_DIR", str(templates))
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:M1")
    sheet['A1'] = '原辅料查验记录'
    sheet.append(['到货日期', '物料名称、型号', '生产日期', '规格', '供应商或生产商', '地址', '联系方式', '保质期', '接收数量', '随车证件', '感官检验', '抽检项目', '备注'])
    workbook.save(templates / '查验台账.xlsx')
    document = Document()
    document.add_paragraph('原辅料检验报告单')
    document.add_paragraph('{{ material_name }} {{ production_dates }}')
    document.add_paragraph('{% for item in items %}{{ item.name }} {{ item.standard }} {{ item.result }} {{ item.conclusion }}{% endfor %}')
    document.add_picture(str(Path(__file__).parent / 'fixtures/label.png'))
    document.save(templates / '检验报告.docx')
    schema = "test_" + uuid.uuid4().hex
    with psycopg.connect(url, autocommit=True) as connection:
        connection.execute(f'CREATE SCHEMA "{schema}"')
    separator = "&" if "?" in url else "?"
    scoped_url = url.replace("postgresql://", "postgresql+psycopg://", 1)
    scoped_url += f"{separator}options=-csearch_path%3D{schema}"
    app = create_app(scoped_url)
    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        if hasattr(app.state, "engine"):
            app.state.engine.dispose()
        with psycopg.connect(url, autocommit=True) as connection:
            connection.execute(f'DROP SCHEMA "{schema}" CASCADE')


def bootstrap(client):
    response = client.post("/api/setup", json={"username": "admin", "password": "test-password-123"})
    assert response.status_code == 201, response.text
    response = client.post("/api/login", json={"username": "admin", "password": "test-password-123"})
    assert response.status_code == 200, response.text
    return response.json()
