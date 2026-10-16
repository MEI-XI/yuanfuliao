import base64
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.template_conversion import _libreoffice, _office_available

from conftest import bootstrap
from test_records import new_ledger


def test_personal_word_template_applies_once_without_changing_server_default(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'inspection_enabled': True, 'fields': {'material_name': '甘氨酸'}, 'report': {'inspection_date': '2026-10-14'}}).json()
    original = client.get('/api/templates/docx')
    assert original.status_code == 200
    document = Document(BytesIO(original.content))
    document.add_paragraph('个人模板签名位置')
    buffer = BytesIO()
    document.save(buffer)
    body = {'record_ids': [record['id']], 'include_pending': True}
    custom = client.post('/api/exports/docx', json=body | {'template_base64': base64.b64encode(buffer.getvalue()).decode()})
    assert custom.status_code == 200
    assert '个人模板签名位置' in Document(BytesIO(custom.content)).element.xml
    assert '甘氨酸_2026.10.14.docx' in unquote(custom.headers['content-disposition'])
    assert client.get('/api/templates/docx').content == original.content
    default = client.post('/api/exports/docx', json=body)
    assert '个人模板签名位置' not in Document(BytesIO(default.content)).element.xml


def test_personal_excel_template_is_used_without_persisting_and_rejects_wrong_files(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'fields': {'arrival_date': '2027-01-02'}}).json()
    original = client.get('/api/templates/xlsx')
    assert original.status_code == 200
    workbook = load_workbook(BytesIO(original.content))
    workbook.active['A1'] = '我的原辅料查验记录X年'
    output = BytesIO()
    workbook.save(output)
    body = {'record_ids': [record['id']], 'include_pending': True}
    custom = client.post('/api/exports/xlsx', json=body | {'template_base64': base64.b64encode(output.getvalue()).decode()})
    assert custom.status_code == 200
    assert load_workbook(BytesIO(custom.content)).active['A1'].value == '我的原辅料查验记录2027年'
    assert client.get('/api/templates/xlsx').content == original.content
    invalid = client.post('/api/exports/xlsx', json=body | {'template_base64': base64.b64encode(b'not an office file').decode()})
    assert invalid.status_code == 422
    assert '模板' in invalid.json()['detail']


def test_broken_or_unsafe_word_template_returns_actionable_error(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'inspection_enabled': True}).json()
    body = {'record_ids': [record['id']], 'include_pending': True}
    broken = BytesIO()
    with ZipFile(broken, 'w') as archive:
        archive.writestr('word/document.xml', '<document/>')
    response = client.post('/api/exports/docx', json=body | {'template_base64': base64.b64encode(broken.getvalue()).decode()})
    assert response.status_code == 422
    unsafe = Document()
    unsafe.add_paragraph('{{ cycler.__init__.__globals__.os.name }}')
    output = BytesIO()
    unsafe.save(output)
    rejected = client.post('/api/exports/docx', json=body | {'template_base64': base64.b64encode(output.getvalue()).decode()})
    assert rejected.status_code == 422
    assert '模板' in rejected.json()['detail']


def test_word_package_with_invalid_document_structure_is_rejected(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'inspection_enabled': True}).json()
    original = client.get('/api/templates/docx').content
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, 'w') as target:
        for name in source.namelist():
            target.writestr(name, '<document/>' if name == 'word/document.xml' else source.read(name))
    response = client.post('/api/exports/docx', json={'record_ids': [record['id']], 'include_pending': True, 'template_base64': base64.b64encode(output.getvalue()).decode()})
    assert response.status_code == 422


def test_plain_report_form_fills_fields_and_all_inspection_rows_without_placeholders(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={
        'inspection_enabled': True, 'fields': {'material_name': '甘氨酸', 'production_dates': ['2026-01-14']},
        'report': {'inspection_date': '2026-10-15', 'conclusion': '符合要求'},
        'items': [{'name': f'项目{i}', 'standard': '≤1', 'result': '0.5', 'conclusion': '符合'} for i in range(6)],
    }).json()
    document = Document()
    document.add_paragraph('原辅料检验报告单')
    table = document.add_table(rows=7, cols=4)
    table.cell(0, 0).text = '检品名称'
    table.cell(0, 2).text = '供应商'
    table.cell(6, 0).text = '检验项目            标准规定                   检验结果                 判定结论'
    for label in ['感官：', '气味：', '滋味：', '水分 %：', '结  论：', '检验员：        批准：']:
        table.cell(6, 0).add_paragraph(label)
    data = BytesIO(); document.save(data)
    response = client.post('/api/exports/docx', json={'record_ids': [record['id']], 'include_pending': True,
        'template_base64': base64.b64encode(data.getvalue()).decode()})
    assert response.status_code == 200
    result = Document(BytesIO(response.content))
    assert result.tables[0].cell(0, 1).text == '甘氨酸'
    assert result.tables[0].cell(3, 3).text == '2026.01.14'
    assert result.tables[0].cell(4, 3).text == '2026.10.15'
    assert result.tables[0].cell(1, 1).text == '\\'
    assert result.tables[0].cell(5, 2).text == '其他检验项目依据该公司提供的内检报告为准'
    for row in result.tables[0].rows[:6]:
        for cell in row.cells:
            assert cell.vertical_alignment == 1
            assert all(paragraph.alignment == 1 for paragraph in cell.paragraphs)
    inspection_table = result.tables[0].cell(6, 0).tables[0]
    assert [cell.text for cell in inspection_table.rows[0].cells] == ['检验项目', '标准规定', '检验结果', '判定结论']
    for row in inspection_table.rows:
        for cell in row.cells:
            assert cell.paragraphs[0].alignment == 1
            assert not cell._tc.xpath('./w:tcPr/w:noWrap')
    content = result.element.xml
    assert all(f'项目{i}' in content for i in range(6))
    assert '符合要求' in content and '{{' not in content


@pytest.mark.skipif(not (_office_available('docx') or _libreoffice()), reason='Requires installed Word or LibreOffice')
@pytest.mark.parametrize('item_count', [4, 30])
def test_floating_plain_report_keeps_title_on_first_pdf_page(client, item_count):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={
        'inspection_enabled': True, 'fields': {'material_name': '甘氨酸'},
        'items': [{'name': f'第{index}项检验', 'standard': '≤1', 'result': '0.5', 'conclusion': '符合'}
                  for index in range(1, item_count + 1)],
    }).json()
    document = Document()
    table = document.add_table(rows=7, cols=4)
    positioning = OxmlElement('w:tblpPr')
    for attribute, value in [('vertAnchor', 'page'), ('horzAnchor', 'margin'), ('tblpY', '2176')]:
        positioning.set(qn('w:' + attribute), value)
    table._tbl.tblPr.append(positioning)
    table.cell(0, 0).text = '检品名称'
    table.cell(0, 2).text = '供应商'
    inspection = table.cell(6, 0).merge(table.cell(6, 3))
    inspection.text = '检验项目            标准规定                   检验结果                 判定结论'
    for label in ['感官：', '气味：', '滋味：', '水分 %：', '结  论：', '检验员：        批准：']:
        inspection.add_paragraph(label)
    inspection.paragraphs[-1].add_run().add_picture(str(Path(__file__).parent / 'fixtures/label.png'), width=Pt(40))
    # The original plain form stores its title after its page-positioned table.
    document.add_paragraph('原辅料检验报告单')
    template = BytesIO()
    document.save(template)
    body = {'record_ids': [record['id']], 'include_pending': True,
            'template_base64': base64.b64encode(template.getvalue()).decode()}
    exported = client.post('/api/exports/docx', json=body)
    assert exported.status_code == 200, exported.text
    with ZipFile(BytesIO(template.getvalue())) as source, ZipFile(BytesIO(exported.content)) as result:
        images = {name: source.read(name) for name in source.namelist() if name.startswith('word/media/')}
        assert images
        assert {name: result.read(name) for name in images} == images
    preview = client.post('/api/previews/docx', json=body)
    assert preview.status_code == 200, preview.text
    pdf = PdfReader(BytesIO(preview.content))
    assert len(pdf.pages) == 1 if item_count == 4 else len(pdf.pages) > 1
    page_texts = [page.extract_text().replace(' ', '') for page in pdf.pages]
    assert '原辅料检验报告单' in page_texts[0]
    assert ''.join(page_texts).count('原辅料检验报告单') == 1
    assert f'第{item_count}项检验' in ''.join(page_texts)
