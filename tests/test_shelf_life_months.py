import base64
import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Material, Record
from conftest import bootstrap
from test_recognition import run_label
from test_records import new_ledger


@pytest.mark.parametrize('value, expected, expiry', [
    ('2年', '24个月', '2028-01-13'),
    ('12月', '12个月', '2027-01-13'),
    ('12', '12个月', '2027-01-13'),
    ('一年', '12个月', '2027-01-13'),
    ('二年', '24个月', '2028-01-13'),
    ('三年', '36个月', '2029-01-13'),
    ('1年6个月', '18个月', '2027-07-13'),
])
def test_derive_returns_months_with_or_without_production_date(client, value, expected, expiry):
    bootstrap(client)
    without_date = client.post('/api/shelf-life/derive', json={'shelf_life': value})
    assert without_date.status_code == 200, without_date.text
    assert without_date.json()['shelf_life'] == expected
    assert without_date.json()['expiry_date'] is None
    with_date = client.post('/api/shelf-life/derive', json={'production_dates': ['2026-01-14'], 'shelf_life': value})
    assert with_date.status_code == 200, with_date.text
    assert with_date.json()['shelf_life'] == expected
    assert with_date.json()['expiry_date'] == expiry


def test_record_and_material_save_months_without_replacing_explicit_expiry_or_clear(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'fields': {
        'shelf_life': '2年', 'expiry_date': '2028-01-14',
    }}).json()
    material = client.post('/api/materials', json={'name': '甘氨酸', 'shelf_life': '三年'}).json()
    assert record['fields']['shelf_life'] == '24个月'
    assert record['fields']['expiry_date'] == '2028-01-14'
    assert material['shelf_life'] == '36个月'
    with Session(client.app.state.engine) as session:
        assert session.get(Record, record['id']).fields['shelf_life'] == '24个月'
        assert session.get(Material, material['id']).shelf_life == '36个月'
    changed = client.patch(f'/api/records/{record["id"]}', json={
        'version': record['version'], 'fields': {'shelf_life': '1年6个月'},
    }).json()
    assert changed['fields']['shelf_life'] == '18个月'
    assert changed['fields']['expiry_date'] == '2028-01-14'
    assert client.put(f'/api/materials/{material["id"]}', json={'name': '甘氨酸', 'shelf_life': '12月'}).json()['shelf_life'] == '12个月'
    cleared = client.patch(f'/api/records/{record["id"]}', json={
        'version': changed['version'], 'fields': {'shelf_life': '', 'expiry_date': None},
    }).json()
    assert cleared['fields']['shelf_life'] == ''
    assert cleared['fields']['expiry_date'] is None


@pytest.mark.parametrize('value', ['90天', '常温12个月，冷藏18个月', ''])
def test_days_ambiguous_values_and_blanks_stay_unchanged(client, value):
    bootstrap(client)
    ledger = new_ledger(client)
    result = client.post('/api/shelf-life/derive', json={'shelf_life': value}).json()
    assert result['shelf_life'] == value
    record = client.post(f'/api/ledgers/{ledger}/records', json={'fields': {'shelf_life': value}}).json()
    material = client.post('/api/materials', json={'name': '甘氨酸', 'shelf_life': value}).json()
    assert record['fields']['shelf_life'] == value
    assert material['shelf_life'] == value


def test_day_duration_keeps_its_expiry_calculation_and_requests_manual_review(client):
    bootstrap(client)
    result = client.post('/api/shelf-life/derive', json={
        'production_dates': ['2026-01-14'], 'shelf_life': '90天',
    }).json()
    assert result['shelf_life'] == '90天'
    assert result['expiry_date'] == '2026-04-13'
    assert '保质期以天数标注，已保留原文，请人工核对' in result['warnings']


def test_legacy_record_and_material_display_months_without_rewriting_database(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={}).json()
    material = client.post('/api/materials', json={'name': '甘氨酸'}).json()
    with Session(client.app.state.engine) as session:
        stored = session.get(Record, record['id'])
        stored.fields = stored.fields | {'shelf_life': '2年'}
        session.get(Material, material['id']).shelf_life = '12月'
        session.commit()
    assert client.get(f'/api/records/{record["id"]}').json()['fields']['shelf_life'] == '24个月'
    assert client.get(f'/api/ledgers/{ledger}/records').json()[0]['fields']['shelf_life'] == '24个月'
    assert client.get('/api/materials').json()[0]['shelf_life'] == '12个月'
    with Session(client.app.state.engine) as session:
        assert session.get(Record, record['id']).fields['shelf_life'] == '2年'
        assert session.get(Material, material['id']).shelf_life == '12月'


@pytest.mark.parametrize('kind', ['xlsx', 'docx'])
def test_legacy_shelf_life_exports_in_months(client, kind):
    bootstrap(client)
    record = client.post(f'/api/ledgers/{new_ledger(client)}/records', json={'inspection_enabled': True}).json()
    with Session(client.app.state.engine) as session:
        stored = session.get(Record, record['id'])
        stored.fields = stored.fields | {'shelf_life': '2年'}
        session.commit()
    body = {'record_ids': [record['id']], 'include_pending': True}
    if kind == 'docx':
        template = Document()
        template.add_paragraph('保质期：{{ shelf_life }}')
        buffer = BytesIO()
        template.save(buffer)
        body['template_base64'] = base64.b64encode(buffer.getvalue()).decode()
    response = client.post(f'/api/exports/{kind}', json=body)
    assert response.status_code == 200, response.text
    if kind == 'xlsx':
        assert load_workbook(BytesIO(response.content)).active['H3'].value == '24个月'
    else:
        assert Document(BytesIO(response.content)).paragraphs[0].text == '保质期：24个月'


@pytest.mark.parametrize('printed, expected', [
    (['保质期：2年'], '24个月'),
    (['保质期：三年'], '36个月'),
    (['保质期：1年6个月'], '18个月'),
    (['保质期：2年', '保质期：24个月'], '24个月'),
])
def test_local_ocr_normalizes_years_and_does_not_conflict_with_equivalent_months(client, printed, expected):
    bootstrap(client)
    job = run_label(client, new_ledger(client), ['生产日期：2026.01.14', *printed])
    assert job['fields']['shelf_life'] == expected
    assert 'shelf_life' not in job['conflicts']


def test_cloud_result_normalizes_explicit_years_to_months(client, monkeypatch):
    bootstrap(client)
    ledger = new_ledger(client)
    client.put('/api/settings/cloud', json={'api_key': 'test-key'})
    real_client = httpx.Client
    cloud_fields = {'production_dates': ['2026-01-14'], 'shelf_life': '2年'}
    monkeypatch.setattr(httpx, 'Client', lambda **kwargs: real_client(transport=httpx.MockTransport(
        lambda request: httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(cloud_fields)}}]})
    ), **kwargs))
    job = run_label(client, ledger, ['品名：甘氨酸'], cloud_enabled=True)
    assert job['cloud_status'] == 'done', job['cloud_error']
    assert job['cloud_fields']['shelf_life'] == '24个月'
    assert job['cloud_fields']['expiry_date'] == '2028-01-13'
