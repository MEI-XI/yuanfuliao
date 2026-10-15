import base64
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.template_conversion import _libreoffice, _office_available
from conftest import bootstrap
from test_records import new_ledger


def _record(client, ledger, name, inspected=True):
    response = client.post(f'/api/ledgers/{ledger}/records', json={
        'inspection_enabled': inspected,
        'fields': {'material_name': name, 'production_dates': ['2026-01-14'], 'notes': '最右列备注'},
        'items': [{'name': '水分', 'result': '0.1%'}] if inspected else None,
    })
    assert response.status_code == 201, response.text
    return response.json()


@pytest.mark.skipif(not (_office_available('docx') or _libreoffice()), reason='Requires installed Word or LibreOffice')
def test_word_preview_merges_personal_template_reports_in_selected_order(client, tmp_path):
    bootstrap(client)
    ledger = new_ledger(client)
    first = _record(client, ledger, '甘氨酸')
    second = _record(client, ledger, '精氨酸')
    document = Document(BytesIO(client.get('/api/templates/docx').content))
    document.add_paragraph('个人模板预览标记')
    template = BytesIO()
    document.save(template)
    response = client.post('/api/previews/docx', json={
        'record_ids': [second['id'], first['id']], 'include_pending': True,
        'template_base64': base64.b64encode(template.getvalue()).decode(),
    })
    assert response.status_code == 200, response.text
    (tmp_path / 'word-preview.pdf').write_bytes(response.content)
    assert response.headers['content-type'] == 'application/pdf'
    assert response.headers['content-disposition'].startswith('inline;')
    assert 'no-store' in response.headers['cache-control']
    pdf = PdfReader(BytesIO(response.content))
    assert len(pdf.pages) == 2
    for page, name in zip(pdf.pages, ['精氨酸', '甘氨酸']):
        text = page.extract_text()
        assert name in text
        assert '个人模板预览标记' in text
        assert '2026.01.14' in text
        assert '0.1%' in text
        assert len(page.images) > 0
    events = client.get('/api/audit-logs').json()['items']
    preview = next(event for event in events if event['target'] == '/api/previews/docx')
    assert preview['action'] == '预览导出'
    assert preview['details']['record_ids'] == [second['id'], first['id']]


@pytest.mark.skipif(not (_office_available('docx') or _libreoffice()), reason='Requires installed Word or LibreOffice')
def test_single_word_preview_uses_default_template(client):
    bootstrap(client)
    record = _record(client, new_ledger(client), '甘氨酸')
    response = client.post('/api/previews/docx', json={'record_ids': [record['id']], 'include_pending': True})
    assert response.status_code == 200, response.text
    pdf = PdfReader(BytesIO(response.content))
    assert len(pdf.pages) == 1
    assert '甘氨酸' in pdf.pages[0].extract_text()


@pytest.mark.skipif(not (_office_available('xlsx') or _libreoffice()), reason='Requires installed Excel or LibreOffice')
def test_excel_preview_keeps_all_thirteen_columns_on_one_page_per_small_ledger(client, tmp_path):
    bootstrap(client)
    first = _record(client, new_ledger(client, '甲公司'), '甘氨酸', inspected=False)
    second = _record(client, new_ledger(client, '乙公司'), '精氨酸', inspected=False)
    workbook = load_workbook(BytesIO(client.get('/api/templates/xlsx').content))
    workbook.active['M2'] = '个人备注栏'
    template = BytesIO()
    workbook.save(template)
    response = client.post('/api/previews/xlsx', json={
        'record_ids': [first['id'], second['id']], 'include_pending': True,
        'template_base64': base64.b64encode(template.getvalue()).decode(),
    })
    assert response.status_code == 200, response.text
    (tmp_path / 'excel-preview.pdf').write_bytes(response.content)
    pdf = PdfReader(BytesIO(response.content))
    assert len(pdf.pages) == 2
    for page, name in zip(pdf.pages, ['甘氨酸', '精氨酸']):
        text = page.extract_text()
        assert name in text
        assert '个人备注栏' in text
        assert '最右列备注' in text
        assert page.mediabox.width > page.mediabox.height


def test_preview_uses_export_permissions_and_record_eligibility(client):
    assert client.post('/api/previews/docx', json={'record_ids': [1]}).status_code == 401
    bootstrap(client)
    record = _record(client, new_ledger(client), '未检验物料', inspected=False)
    assert client.post('/api/previews/missing', json={'record_ids': [record['id']]}).status_code == 404
    assert client.post('/api/previews/xlsx', json={'record_ids': []}).status_code == 422
    pending = client.post('/api/previews/xlsx', json={'record_ids': [record['id']]})
    assert pending.status_code == 409
    assert '尚未核对' in pending.json()['detail']
    uninspected = client.post('/api/previews/docx', json={'record_ids': [record['id']], 'include_pending': True})
    assert uninspected.status_code == 409
    assert '未检验' in uninspected.json()['detail']
    client.post(f'/api/records/{record["id"]}/void', json={'version': record['version']})
    voided = client.post('/api/previews/xlsx', json={'record_ids': [record['id']], 'include_pending': True})
    assert voided.status_code == 409
    assert '作废' in voided.json()['detail']
    assert client.post('/api/previews/docx', json={'record_ids': [999999], 'include_pending': True}).status_code == 404


def test_preview_reports_invalid_template_and_missing_converter(client, monkeypatch):
    bootstrap(client)
    record = _record(client, new_ledger(client), '甘氨酸')
    bad_template = Document()
    bad_template.add_paragraph('{% if %}')
    template = BytesIO()
    bad_template.save(template)
    invalid = client.post('/api/previews/docx', json={
        'record_ids': [record['id']], 'include_pending': True,
        'template_base64': base64.b64encode(template.getvalue()).decode(),
    })
    assert invalid.status_code == 422
    assert '模板' in invalid.json()['detail']
    monkeypatch.setattr('app.template_conversion._libreoffice', lambda: None)
    monkeypatch.setattr('app.template_conversion._office_available', lambda kind: False)
    unavailable = client.post('/api/previews/docx', json={'record_ids': [record['id']], 'include_pending': True})
    assert unavailable.status_code == 503
    assert 'Office' in unavailable.json()['detail']
