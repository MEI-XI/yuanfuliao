from io import BytesIO
from pathlib import Path
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import update
from sqlalchemy.orm import Session

from app.models import Record
from app.report_attachments import ReportAttachment, report_attachment_path
from app.template_conversion import _libreoffice, _office_available

from conftest import bootstrap
from test_records import new_ledger


PDF = b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\n%%EOF\n'


def create_record(client):
    bootstrap(client)
    ledger = new_ledger(client)
    return client.post(f'/api/ledgers/{ledger}/records', json={'fields': {'notes': '保留原记录内容'}}).json()


def upload(client, record, data=PDF, filename='供应商报告.pdf', version=None):
    return client.post(f'/api/records/{record["id"]}/reports',
                       data={'version': record['version'] if version is None else version},
                       files={'file': (filename, data, 'application/octet-stream')})


def test_upload_report_preserves_record_fields_and_resets_checked_state(client):
    record = create_record(client)
    checked = client.post(f'/api/records/{record["id"]}/check', json={'version': record['version']}).json()
    response = upload(client, checked)
    assert response.status_code == 201, response.text
    result = response.json()
    attachment = result['attachment']
    assert attachment['filename'] == '供应商报告.pdf'
    assert attachment['media_type'] == 'application/pdf'
    assert attachment['size'] == len(PDF)
    assert result['record_version'] == 3
    assert result['status'] == 'pending'
    assert result['checker_id'] is None
    assert result['checked_at'] is None
    saved = client.get(f'/api/records/{record["id"]}').json()
    assert saved['fields']['notes'] == '保留原记录内容'
    assert saved['version'] == 3
    assert client.get(f'/api/records/{record["id"]}/reports').json() == [attachment]
    history = client.get(f'/api/records/{record["id"]}/history').json()
    assert history[-1]['action'] == 'report_uploaded'
    file = client.get(f'/api/reports/{attachment["id"]}/file')
    assert file.content == PDF
    assert file.headers['content-type'] == 'application/pdf'
    assert file.headers['content-disposition'].startswith('inline;')
    assert file.headers['x-content-type-options'] == 'nosniff'
    assert client.get(f'/api/reports/{attachment["id"]}/preview').content == PDF


def test_reports_are_readable_to_other_users_but_only_creator_or_admin_can_modify(client):
    record = create_record(client)
    attachment = upload(client, record).json()['attachment']
    client.post('/api/users', json={'username': 'reader', 'password': 'reader-password'})
    client.post('/api/login', json={'username': 'reader', 'password': 'reader-password'})
    assert client.get(f'/api/records/{record["id"]}/reports').status_code == 200
    assert client.get(f'/api/reports/{attachment["id"]}/file').content == PDF
    assert upload(client, record, version=2).status_code == 403
    assert client.delete(f'/api/reports/{attachment["id"]}?version=2').status_code == 403
    client.post('/api/logout')
    assert client.get(f'/api/reports/{attachment["id"]}/file').status_code == 401


def test_report_changes_require_current_version_and_delete_removes_file(client):
    record = create_record(client)
    attachment = upload(client, record).json()['attachment']
    assert upload(client, record).status_code == 409
    assert client.delete(f'/api/reports/{attachment["id"]}?version=1').status_code == 409
    response = client.delete(f'/api/reports/{attachment["id"]}?version=2')
    assert response.status_code == 200, response.text
    assert response.json()['record_version'] == 3
    assert client.get(f'/api/records/{record["id"]}/reports').json() == []
    assert client.get(f'/api/reports/{attachment["id"]}/file').status_code == 404
    import os
    assert list((Path(os.environ['DATA_DIR']) / 'report-attachments').glob('*')) == []


@pytest.mark.parametrize('filename,content', [
    ('report.exe', b'MZpayload'), ('report.pdf', b'MZpayload'), ('report.png', PDF),
    ('report.docx', b'PKnot an office document'), ('report.txt', b'\x00binary'),
    ('empty.pdf', b''),
])
def test_report_upload_rejects_disallowed_or_mismatched_files(client, filename, content):
    record = create_record(client)
    response = upload(client, record, content, filename)
    assert response.status_code == 422, response.text
    assert client.get(f'/api/records/{record["id"]}').json()['version'] == 1


def test_report_upload_rejects_files_larger_than_twenty_megabytes(client):
    record = create_record(client)
    response = upload(client, record, b'%PDF-' + b'0' * (20 * 1024 * 1024), 'oversized.pdf')
    assert response.status_code == 413, response.text


@pytest.mark.parametrize('kind', ['png', 'jpg', 'webp', 'docx', 'xlsx', 'txt'])
def test_common_report_formats_can_be_uploaded_and_downloaded(client, kind):
    record = create_record(client)
    output = BytesIO()
    if kind in ('png', 'jpg', 'webp'):
        Image.new('RGB', (10, 10), 'white').save(output, format={'jpg': 'JPEG'}.get(kind, kind))
    elif kind == 'docx':
        document = Document()
        document.add_paragraph('出厂检验合格报告')
        document.save(output)
    elif kind == 'xlsx':
        workbook = Workbook()
        workbook.active['A1'] = '出厂检验合格报告'
        workbook.save(output)
    else:
        output.write('出厂检验合格报告'.encode())
    data = output.getvalue()
    response = upload(client, record, data, f'检验报告.{kind}')
    assert response.status_code == 201, response.text
    attachment = response.json()['attachment']
    file = client.get(f'/api/reports/{attachment["id"]}/file')
    assert file.content == data
    if kind in ('docx', 'xlsx', 'txt'):
        assert file.headers['content-disposition'].startswith('attachment;')
    else:
        assert file.headers['content-disposition'].startswith('inline;')
    if kind not in ('docx', 'xlsx'):
        preview = client.get(f'/api/reports/{attachment["id"]}/preview')
        assert preview.status_code == 200
        if kind == 'txt':
            assert preview.text == '出厂检验合格报告'
        else:
            assert preview.content == data


@pytest.mark.parametrize('kind', ['docx', 'xlsx'])
def test_office_report_attachment_preview_renders_the_uploaded_document(client, kind):
    if not (_office_available(kind) or _libreoffice()):
        pytest.skip('Requires installed Office or LibreOffice')
    record = create_record(client)
    output = BytesIO()
    if kind == 'docx':
        document = Document()
        document.add_paragraph('供应商原始报告内容')
        document.save(output)
    else:
        workbook = Workbook()
        workbook.active['A1'] = '供应商原始报告内容'
        workbook.save(output)
    attachment = upload(client, record, output.getvalue(), f'报告.{kind}').json()['attachment']
    preview = client.get(f'/api/reports/{attachment["id"]}/preview')
    assert preview.status_code == 200, preview.text
    assert preview.headers['content-type'] == 'application/pdf'
    assert preview.headers['content-disposition'].startswith('inline;')
    assert '供应商原始报告内容' in PdfReader(BytesIO(preview.content)).pages[0].extract_text()


def test_void_reports_stay_readable_but_cannot_be_changed(client):
    record = create_record(client)
    attachment = upload(client, record).json()['attachment']
    assert client.post(f'/api/records/{record["id"]}/void', json={'version': 2}).status_code == 200
    assert client.get(f'/api/reports/{attachment["id"]}/file').content == PDF
    assert upload(client, record, version=3).status_code == 409
    assert client.delete(f'/api/reports/{attachment["id"]}?version=3').status_code == 409


def test_restoring_a_record_preserves_uploaded_report_files(client):
    record = create_record(client)
    attachment = upload(client, record).json()['attachment']
    voided = client.post(f'/api/records/{record["id"]}/void', json={'version': 2}).json()
    response = client.post(f'/api/records/{record["id"]}/restore', json={'version': voided['version']})
    assert response.status_code == 200, response.text
    assert client.get(f'/api/records/{record["id"]}/reports').json() == [attachment]
    assert client.get(f'/api/reports/{attachment["id"]}/file').content == PDF


def test_expired_record_reports_become_unavailable_and_cleanup_deletes_the_files(client):
    record = create_record(client)
    attachment = upload(client, record).json()['attachment']
    client.post(f'/api/records/{record["id"]}/void', json={'version': 2})
    with Session(client.app.state.engine) as session:
        path = report_attachment_path(session.get(ReportAttachment, attachment['id']).storage_key)
        session.execute(update(Record).where(Record.id == record['id']).values(voided_at=datetime.now(timezone.utc) - timedelta(days=31)))
        session.commit()
    assert path.is_file()
    assert client.get(f'/api/reports/{attachment["id"]}/file').status_code == 404
    assert client.get(f'/api/records/{record["id"]}/reports').status_code == 404
    assert client.get('/api/recycle-bin').status_code == 200
    assert not path.exists()
    with Session(client.app.state.engine) as session:
        assert session.get(ReportAttachment, attachment['id']) is None
