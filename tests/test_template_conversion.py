from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from pathlib import Path
from zipfile import ZipFile
import os
import shutil

import pytest

from app.template_conversion import _office_available

from conftest import bootstrap


def test_conversion_accepts_arbitrary_word_filename_without_persisting(client):
    bootstrap(client)
    original = client.get('/api/templates/docx').content
    response = client.post('/api/templates/convert/docx', files={'file': ('任意报告名称.docx', original)})
    assert response.status_code == 200, response.text
    assert response.content == original
    assert 'wordprocessingml.document' in response.headers['content-type']
    assert Document(BytesIO(response.content)).inline_shapes
    assert client.get('/api/templates/docx').content == original


def test_conversion_validates_excel_structure_and_file_type(client):
    bootstrap(client)
    original = client.get('/api/templates/xlsx').content
    response = client.post('/api/templates/convert/xlsx', files={'file': ('公司自用.xlsx', original)})
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content)).active['A1'].value == '原辅料查验记录'
    for filename, content in [('report.docx', original), ('image.png', original), ('坏文件.xlsx', b'bad')]:
        rejected = client.post('/api/templates/convert/xlsx', files={'file': (filename, content)})
        assert rejected.status_code == 422
    workbook = load_workbook(BytesIO(original))
    workbook.active['A1'] = '错误标题'
    output = BytesIO()
    workbook.save(output)
    rejected = client.post('/api/templates/convert/xlsx', files={'file': ('自用.xlsx', output.getvalue())})
    assert rejected.status_code == 422


@pytest.mark.parametrize('kind,extension,content_type', [
    ('docx', 'dotm', 'application/vnd.ms-word.template.macroEnabledTemplate.main+xml'),
    ('xlsx', 'xltm', 'application/vnd.ms-excel.template.macroEnabled.main+xml'),
])
def test_conversion_normalizes_office_template_packages_and_preserves_media(client, kind, extension, content_type):
    bootstrap(client)
    original = client.get(f'/api/templates/{kind}').content
    if kind == 'xlsx':
        workbook = load_workbook(BytesIO(original))
        workbook.active.add_image(Image(str(Path(__file__).parent / 'fixtures/label.png')), 'A3')
        workbook_bytes = BytesIO()
        workbook.save(workbook_bytes)
        original = workbook_bytes.getvalue()
    output = BytesIO()
    standard_type = ('application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'
                     if kind == 'docx' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml')
    with ZipFile(BytesIO(original)) as source, ZipFile(output, 'w') as target:
        for name in source.namelist():
            content = source.read(name)
            if name == '[Content_Types].xml':
                content = content.replace(standard_type.encode(), content_type.encode())
            target.writestr(name, content)
    response = client.post(f'/api/templates/convert/{kind}', files={'file': ('我的模板.' + extension, output.getvalue())})
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as normalized:
        assert standard_type.encode() in normalized.read('[Content_Types].xml')
        if kind == 'docx':
            assert any(name.startswith('word/media/') for name in normalized.namelist())
        else:
            assert any(name.startswith('xl/media/') for name in normalized.namelist())


def test_conversion_requires_login(client):
    response = client.post('/api/templates/convert/docx', files={'file': ('模板.docx', b'bad')})
    assert response.status_code == 401


@pytest.mark.skipif(not (_office_available('docx') or shutil.which('soffice') or os.getenv('LIBREOFFICE_PATH')), reason='Requires installed Word or LibreOffice')
def test_conversion_of_real_rtf_preserves_document_text(client):
    bootstrap(client)
    response = client.post('/api/templates/convert/docx', files={'file': ('随意命名.rtf', br'{\rtf1\ansi Original report \{\{ material_name \}\}}')})
    assert response.status_code == 200, response.text
    assert 'Original report {{ material_name }}' in Document(BytesIO(response.content)).element.xml


@pytest.mark.skipif(not (_office_available('xlsx') or shutil.which('soffice') or os.getenv('LIBREOFFICE_PATH')), reason='Requires installed Excel or LibreOffice')
def test_conversion_of_real_csv_validates_normalized_ledger(client):
    bootstrap(client)
    csv = ('原辅料查验记录,,,,,,,,,,,,\r\n' + ','.join('字段' + str(number) for number in range(13)) + '\r\n').encode('utf-8-sig')
    response = client.post('/api/templates/convert/xlsx', files={'file': ('自己的台账.csv', csv)})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active['A1'].value == '原辅料查验记录'
    assert workbook.active.max_column == 13
