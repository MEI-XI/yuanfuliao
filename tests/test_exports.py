from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from docx import Document
from sqlalchemy.orm import Session
from zipfile import ZipFile

from conftest import bootstrap
from test_records import new_ledger


def test_word_long_product_name_fits_windows_filename_limit(client):
    bootstrap(client)
    ledger = new_ledger(client)
    body = {'fields': {'material_name': '原料' * 140}, 'report': {'inspection_date': '2026-10-14'}}
    first = client.post(f'/api/ledgers/{ledger}/records', json=body).json()
    second = client.post(f'/api/ledgers/{ledger}/records', json=body).json()
    response = client.post('/api/exports/docx', json={'record_ids': [first['id'], second['id']], 'include_pending': True})
    with ZipFile(BytesIO(response.content)) as archive:
        assert all(len(name.encode('utf-16-le')) // 2 <= 255 for name in archive.namelist())
        assert len(set(archive.namelist())) == 2


def test_excel_preserves_13_columns_company_sheets_blanks_and_manufacturer_source(client):
    bootstrap(client)
    first_ledger = new_ledger(client)
    second_ledger = new_ledger(client, "味源食品")
    first = client.post(f"/api/ledgers/{first_ledger}/records", json={"fields": {
        "arrival_date": None, "material_name": "甘氨酸", "company_name": "甲商行",
        "address": "乙工厂地址", "address_source": "manufacturer", "contact": "0311-12345678",
        "contact_source": "manufacturer", "production_dates": ["2026-01-14", "2026-01-15"]
    }})
    assert first.status_code == 201
    second = client.post(f"/api/ledgers/{second_ledger}/records", json={}).json()
    ids = [first.json()["id"], second["id"]]
    assert client.post("/api/exports/xlsx", json={"record_ids": ids}).status_code == 409
    response = client.post("/api/exports/xlsx", json={"record_ids": ids, "include_pending": True})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["青山科技", "味源食品"]
    sheet = workbook["青山科技"]
    assert sheet.max_column == 13
    assert [sheet.cell(3, col).value for col in range(1, 10)] == [
        None, "甘氨酸", "2026.01.14\n2026.01.15", None, "甲商行", "生产商：乙工厂地址", "生产商：0311-12345678", None, None
    ]
    client.post(f'/api/records/{ids[0]}/void', json={"version": 1})
    assert client.post("/api/exports/xlsx", json={"record_ids": ids, "include_pending": True}).status_code == 409


def test_word_uses_dotted_production_date_one_report_per_record_and_current_template_signatures(client):
    bootstrap(client)
    ledger = new_ledger(client)
    saved = client.post(f"/api/ledgers/{ledger}/records", json={"fields": {
        "material_name": "甘氨酸", "production_dates": ["2026-01-14"]},
        "items": [{"name": "水分", "standard": "≤0.2%", "result": "0.1%", "conclusion": "符合"}]
    }).json()
    response = client.post("/api/exports/docx", json={"record_ids": [saved["id"]], "include_pending": True})
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    content = document.element.xml
    assert "2026.01.14" in content and "甘氨酸" in content
    assert "0.1%" in content and "≤0.2%" in content
    assert "{{" not in content
    from app.exporting import template_dir
    with ZipFile(template_dir() / "检验报告.docx") as template_zip, ZipFile(BytesIO(response.content)) as output_zip:
        pictures = [name for name in template_zip.namelist() if name.startswith("word/media/")]
        assert pictures
        assert {name: template_zip.read(name) for name in pictures} == {name: output_zip.read(name) for name in pictures}
    second = client.post(f"/api/ledgers/{ledger}/records", json={}).json()
    batch = client.post("/api/exports/docx", json={"record_ids": [saved["id"], second["id"]], "include_pending": True})
    with ZipFile(BytesIO(batch.content)) as archive:
        assert len(archive.namelist()) == 2
        assert all(name.endswith(".docx") for name in archive.namelist())


def test_company_named_like_template_exports_exact_sheet_name(client):
    bootstrap(client)
    ledger = new_ledger(client, 'Sheet')
    record = client.post(f'/api/ledgers/{ledger}/records', json={}).json()
    response = client.post('/api/exports/xlsx', json={'record_ids': [record['id']], 'include_pending': True})
    assert load_workbook(BytesIO(response.content)).sheetnames == ['Sheet']


def test_word_preserves_inspection_unit_in_existing_project_column(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'items': [{'name': '水分', 'standard': '≤0.2', 'unit': 'mg/kg', 'result': '0.1'}]}).json()
    response = client.post('/api/exports/docx', json={'record_ids': [record['id']], 'include_pending': True})
    assert '水分（mg/kg）' in Document(BytesIO(response.content)).element.xml


@pytest.mark.parametrize('template_title', ['原辅料查验记录', '原辅料查验记录X年', '原辅料查验记录2024年'])
def test_excel_title_uses_first_selected_arrival_year_in_each_company(client, template_title):
    from app.exporting import template_dir

    bootstrap(client)
    ledgers = [new_ledger(client, name) for name in ['甲公司', '乙公司', '丙公司']]
    records = []
    for ledger, arrivals in zip(ledgers, [('2025-12-31', '2026-01-01'), ('2028-02-03',), (None, '2029-04-05')]):
        records.append([client.post(f'/api/ledgers/{ledger}/records', json={'fields': {
            'arrival_date': arrival, 'material_name': f'产品{index}'
        }}).json()['id'] for index, arrival in enumerate(arrivals)])
    workbook = load_workbook(template_dir() / '查验台账.xlsx')
    workbook.active['A1'] = template_title
    workbook.save(template_dir() / '查验台账.xlsx')

    response = client.post('/api/exports/xlsx', json={
        'record_ids': [records[0][1], records[1][0], records[0][0], *records[2]], 'include_pending': True
    })
    assert response.status_code == 200
    exported = load_workbook(BytesIO(response.content))
    assert exported['甲公司']['A1'].value == '原辅料查验记录2026年'
    assert exported['乙公司']['A1'].value == '原辅料查验记录2028年'
    assert exported['丙公司']['A1'].value == '原辅料查验记录'
    assert exported['甲公司']['A3'].value.year == 2026
    assert exported['甲公司']['A4'].value.year == 2025


def test_word_zip_uses_inspection_date_and_shanghai_creation_date_with_unique_safe_names(client):
    from app.models import Record

    bootstrap(client)
    ledger = new_ledger(client)
    records = []
    products = ['甘氨酸', '甘氨酸', '甘氨酸', 'Product', 'product', '盐/糖:*?"<>|\\', '']
    for index, product in enumerate(products):
        response = client.post(f'/api/ledgers/{ledger}/records', json={
            'fields': {'material_name': product, 'arrival_date': '2023-01-01'},
            'report': {'inspection_date': '2026-10-14'} if index < 6 else {},
            'items': [{'name': f'检验项{index}', 'result': f'结果{index}'}]
        })
        assert response.status_code == 201
        records.append(response.json()['id'])
    # Only this test's isolated schema is used; the boundary crosses midnight in Shanghai.
    with Session(client.app.state.engine) as session:
        for record_id in records:
            session.get(Record, record_id).created_at = datetime(2026, 10, 15, 16, 30, tzinfo=timezone.utc)
        session.commit()

    response = client.post('/api/exports/docx', json={'record_ids': records, 'include_pending': True})
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/zip'
    with ZipFile(BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert names == [
            '甘氨酸_2026.10.14.docx', '甘氨酸_2026.10.14(1).docx', '甘氨酸_2026.10.14(2).docx',
            'Product_2026.10.14.docx', 'product_2026.10.14(1).docx',
            '盐_糖_________2026.10.14.docx', '未命名_2026.10.16.docx'
        ]
        assert len({name.casefold() for name in names}) == len(records)
        for index, name in enumerate(names):
            document = Document(BytesIO(archive.read(name)))
            assert f'检验项{index}' in document.element.xml
            assert f'结果{index}' in document.element.xml


def test_excel_uses_supplied_template_bytes(monkeypatch, tmp_path):
    from app.exporting import excel_bytes
    from app.models import Record

    monkeypatch.setenv('TEMPLATE_DIR', str(tmp_path / 'missing'))
    workbook = Workbook()
    workbook.active['A1'] = '原辅料查验记录X年'
    workbook.active['M2'] = '自定义备注栏'
    template = BytesIO()
    workbook.save(template)
    record = Record(ledger_id=1, fields={'arrival_date': '2027-01-01', 'material_name': '甘氨酸'})

    exported = load_workbook(BytesIO(excel_bytes([record], {1: '甲公司'}, template.getvalue())))
    assert exported['甲公司']['A1'].value == '原辅料查验记录2027年'
    assert exported['甲公司']['M2'].value == '自定义备注栏'
    assert exported['甲公司']['B3'].value == '甘氨酸'


def test_word_archive_renders_supplied_template_fresh_for_each_record(monkeypatch, tmp_path):
    from app.exporting import word_archive
    from app.models import Record

    monkeypatch.setenv('TEMPLATE_DIR', str(tmp_path / 'missing'))
    document = Document()
    document.add_paragraph('自定义报告：{{ material_name }}')
    document.add_picture(str(Path(__file__).parent / 'fixtures/label.png'))
    template = BytesIO()
    document.save(template)
    records = [Record(fields={'material_name': name}, report={'inspection_date': '2026-10-14'}, items=[])
               for name in ['甲品', '乙品']]

    with ZipFile(BytesIO(word_archive(records, template.getvalue()))) as archive:
        for product in ['甲品', '乙品']:
            content = archive.read(f'{product}_2026.10.14.docx')
            assert f'自定义报告：{product}' in Document(BytesIO(content)).element.xml
            with ZipFile(BytesIO(content)) as report, ZipFile(BytesIO(template.getvalue())) as original:
                images = [name for name in original.namelist() if name.startswith('word/media/')]
                assert {name: report.read(name) for name in images} == {name: original.read(name) for name in images}


def test_supplied_word_template_cannot_access_python_internals():
    from app.exporting import word_bytes
    from app.models import Record
    from jinja2.exceptions import SecurityError

    document = Document()
    document.add_paragraph("{{ ''.__class__.__mro__[1].__subclasses__() }}")
    template = BytesIO()
    document.save(template)
    record = Record(fields={}, report={}, items=[])
    with pytest.raises(SecurityError):
        word_bytes(record, template.getvalue())
