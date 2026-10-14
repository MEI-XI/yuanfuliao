from datetime import datetime
import base64
from zoneinfo import ZoneInfo
from io import BytesIO

from docx import Document

from conftest import bootstrap
from test_records import new_ledger


def test_new_report_defaults_today_but_saved_date_can_change_or_be_blank(client):
    bootstrap(client)
    ledger = new_ledger(client)
    created = client.post(f'/api/ledgers/{ledger}/records', json={'inspection_enabled': True}).json()
    assert created['report']['inspection_date'] == datetime.now(ZoneInfo('Asia/Shanghai')).date().isoformat()
    changed = client.patch(f'/api/records/{created["id"]}', json={'version': 1, 'report': {'inspection_date': '2026-01-14'}}).json()
    reopened = client.get(f'/api/records/{created["id"]}').json()
    assert reopened['report']['inspection_date'] == '2026-01-14'
    cleared = client.patch(f'/api/records/{created["id"]}', json={'version': changed['version'], 'report': {'inspection_date': None}}).json()
    assert cleared['report']['inspection_date'] is None


def test_not_inspected_records_lock_report_and_block_word_exports_but_allow_excel(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={'fields': {'material_name': '未检验物料'}}).json()
    assert record['inspection_enabled'] is False
    blocked = client.post('/api/exports/docx', json={'record_ids': [record['id']], 'include_pending': True})
    assert blocked.status_code == 409
    assert '未检验' in blocked.json()['detail']
    assert client.post('/api/exports/xlsx', json={'record_ids': [record['id']], 'include_pending': True}).status_code == 200
    invalid = client.patch(f'/api/records/{record["id"]}', json={'version': 1, 'items': [{'name': '水分', 'result': '0.1'}]})
    assert invalid.status_code == 409
    enabled = client.patch(f'/api/records/{record["id"]}', json={'version': 1, 'inspection_enabled': True, 'items': [{'name': '水分', 'result': '0.1'}], 'report': {'inspection_date': '2026-01-14'}}).json()
    assert enabled['inspection_enabled'] is True
    assert client.post('/api/exports/docx', json={'record_ids': [record['id']], 'include_pending': True}).status_code == 200
    disabled = client.patch(f'/api/records/{record["id"]}', json={'version': enabled['version'], 'inspection_enabled': False}).json()
    assert disabled['items'] == enabled['items']
    assert disabled['report'] == enabled['report']
    assert disabled['status'] == 'pending'
    assert client.get('/api/records').json()[0]['inspection_enabled'] is False
    again = client.patch(f'/api/records/{record["id"]}', json={'version': disabled['version'], 'inspection_enabled': True}).json()
    assert again['items'] == enabled['items']
    mixed = client.post(f'/api/ledgers/{ledger}/records', json={}).json()
    assert client.post('/api/exports/docx', json={'record_ids': [record['id'], mixed['id']], 'include_pending': True}).status_code == 409


def test_explicit_draft_save_retains_work_after_toggling_inspection_back_off(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={}).json()
    saved = client.patch(f'/api/records/{record["id"]}', json={
        'version': record['version'], 'inspection_enabled': False,
        'items': [{'name': '水分', 'result': '0.1'}],
        'report': {'inspection_date': '2026-01-14'},
    })
    assert saved.status_code == 200
    assert saved.json()['inspection_enabled'] is False
    restored = client.patch(f'/api/records/{record["id"]}', json={
        'version': saved.json()['version'], 'inspection_enabled': True,
    }).json()
    assert restored['items'][0]['result'] == '0.1'
    assert restored['report']['inspection_date'] == '2026-01-14'


def test_report_uses_template_signature_images_without_legacy_manual_names(client):
    bootstrap(client)
    ledger = new_ledger(client)
    record = client.post(f'/api/ledgers/{ledger}/records', json={
        'inspection_enabled': True,
        'report': {'inspector': '旧检验员姓名', 'approver': '旧批准人姓名'},
    }).json()
    template = Document(BytesIO(client.get('/api/templates/docx').content))
    template.add_paragraph('检验员：{{ inspector }} 批准：{{ approver }}')
    content = BytesIO()
    template.save(content)
    exported = client.post('/api/exports/docx', json={
        'record_ids': [record['id']], 'include_pending': True,
        'template_base64': base64.b64encode(content.getvalue()).decode(),
    })
    assert exported.status_code == 200
    document = Document(BytesIO(exported.content))
    assert '旧检验员姓名' not in document.element.xml
    assert '旧批准人姓名' not in document.element.xml
    assert len(document.inline_shapes) == len(template.inline_shapes) > 0
