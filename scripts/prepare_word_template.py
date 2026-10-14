"""Convert an already-converted DOCX copy into the editable report template.

Usage: python scripts/prepare_word_template.py input.docx output.docx
The original input is never overwritten. Embedded signature images are retained.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def prepare(source, output):
    if Path(source).resolve() == Path(output).resolve() or Path(output).exists():
        raise ValueError("输出必须是不存在的新文件，以保护已修改的模板")
    document = Document(source)
    table = document.tables[0]
    # The DOC has a floating table preceding the title in document order. A long
    # report otherwise sends the title to the last page. Keep both in text flow.
    for position in table._tbl.tblPr.xpath('./w:tblpPr'):
        position.getparent().remove(position)
    title = next(p for p in document.paragraphs if '检验报告' in p.text)
    table._tbl.addprevious(title._p)
    title.paragraph_format.space_after = Pt(40)
    title.paragraph_format.keep_with_next = True
    fields = {(0, 1): "material_name", (0, 3): "company_name", (1, 1): "batch_blank", (1, 3): "specification",
              (2, 1): "packaging", (2, 3): "received_quantity", (3, 1): "basis", (3, 3): "production_dates",
              (4, 1): "shelf_life", (4, 3): "inspection_date"}
    for (row, col), field in fields.items():
        table.cell(row, col).paragraphs[0].add_run("{{ " + field + " }}")
    table.cell(5, 0).text = "该公司是否提供内检报告：{{ internal_report }}"
    table.cell(5, 2).text = "{{ report_reference }}"
    cell = table.cell(6, 0)
    paragraphs = list(cell.paragraphs)
    # The last two paragraphs contain the original floating signature shapes.
    for paragraph in paragraphs[:5]:
        cell._tc.remove(paragraph._p)
    conclusion, signature = paragraphs[5:7]
    for paragraph in (conclusion, signature):
        for node in paragraph._p.xpath(".//w:t"):
            node.text = ""
    conclusion.add_run("结  论：{{ conclusion }}")
    signature.paragraph_format.tab_stops.add_tab_stop(Inches(4.45))
    signature.paragraph_format.space_after = Pt(18)
    signature.add_run("检验员：{{ inspector }}\t批准：{{ approver }}")
    inner = cell.add_table(rows=4, cols=4)
    inner.autofit = False
    for column, width in zip(inner.columns, [1.45, 2.05, 1.6, 1.65]):
        column.width = Inches(width)
    for col, label in enumerate(["检验项目", "标准规定", "检验结果", "判定结论"]):
        inner.cell(0, col).text = label
    inner.cell(1, 0).text = "{%tr for item in items %}"
    for col, key in enumerate(["name", "standard", "result", "conclusion"]):
        inner.cell(2, col).text = "{{ item." + key + " }}"
    inner.cell(3, 0).text = "{%tr endfor %}"
    cell._tc.remove(inner._tbl)
    cell._tc.insert(cell._tc.index(conclusion._p), inner._tbl)
    for row in inner.rows:
        for entry in row.cells:
            for paragraph in entry.paragraphs:
                paragraph.paragraph_format.space_after = Pt(8)
                for run in paragraph.runs:
                    run.font.size = Pt(11)
    document.save(output)


if __name__ == "__main__":
    prepare(*sys.argv[1:3])
